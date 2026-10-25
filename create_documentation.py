import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_placemate_doc():
    doc = Document()

    # 1. Page Margins (1 inch everywhere)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 2. Configure Default Styles
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(11)
    normal_font.color.rgb = RGBColor(30, 30, 30)

    # Helpers
    def set_font(run, font_name="Times New Roman", size_pt=11, bold=False, italic=False, color=(30, 30, 30)):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color)

    def add_title(text, subtitle=""):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        set_font(r, size_pt=18, bold=True, color=(15, 23, 42))

        if subtitle:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_after = Pt(14)
            r_sub = p_sub.add_run(subtitle)
            set_font(r_sub, size_pt=12, italic=True, color=(100, 116, 139))

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        set_font(r, size_pt=14, bold=True, color=(15, 23, 42))
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        set_font(r, size_pt=12, bold=True, color=(30, 41, 59))
        return p

    def add_body_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(5)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            set_font(r_pre, size_pt=11, bold=True)
        r = p.add_run(text)
        set_font(r, size_pt=11, bold=False)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            set_font(r_pre, size_pt=11, bold=True)
        r = p.add_run(text)
        set_font(r, size_pt=11, bold=False)
        return p

    def format_table_header(row, col_widths=None):
        for idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
            tcPr.append(shd)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_font(r, size_pt=11, bold=True, color=(15, 23, 42))
            if col_widths and idx < len(col_widths):
                cell.width = col_widths[idx]

    def format_table_cells(table, col_widths=None):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row_idx, row in enumerate(table.rows):
            if row_idx == 0:
                format_table_header(row, col_widths)
                continue
            for c_idx, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.1
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        set_font(r, size_pt=10, bold=r.bold, color=(51, 65, 85))
                if col_widths and c_idx < len(col_widths):
                    cell.width = col_widths[c_idx]

    def add_screenshot(img_rel_path, figure_no, title, description, width_in=4.8):
        full_path = os.path.join(os.getcwd(), img_rel_path)
        if not os.path.exists(full_path):
            print(f"Warning: image not found: {full_path}")
            return

        # Add image centered
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        
        run_img = p_img.add_run()
        run_img.add_picture(full_path, width=Inches(width_in))

        # Add caption
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(f"Figure {figure_no}: {title}")
        set_font(r_cap, size_pt=10, bold=True, italic=True, color=(71, 85, 105))

        # Add brief simple explanation
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.line_spacing = 1.12
        p_desc.paragraph_format.space_after = Pt(8)
        r_desc = p_desc.add_run(description)
        set_font(r_desc, size_pt=10.5, bold=False, color=(51, 65, 85))

    # =========================================================================
    # FOOTER: DYNAMIC PAGE NUMBERING
    # =========================================================================
    for sec in doc.sections:
        footer = sec.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_foot.text = ""
        r_pre = p_foot.add_run("Page ")
        set_font(r_pre, size_pt=10, color=(100, 116, 139))
        
        r_pg = p_foot.add_run()
        set_font(r_pg, size_pt=10, color=(100, 116, 139))
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        r_pg._r.append(fldSimple)

    # =========================================================================
    # PAGE 1: TITLE AND INDEX
    # =========================================================================
    add_title("PlaceMate: Placement Management Portal", "Project Report and System Documentation")
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(18)
    r_meta = p_meta.add_run("A Unified Web-Based Campus Placement Automation Platform\nAcademic Project Documentation")
    set_font(r_meta, size_pt=11, italic=False, color=(71, 85, 105))

    add_heading_1("Index")
    
    index_table = doc.add_table(rows=1, cols=3)
    hdr_cells = index_table.rows[0].cells
    hdr_cells[0].paragraphs[0].add_run("Sr. No.")
    hdr_cells[1].paragraphs[0].add_run("Topic Name")
    hdr_cells[2].paragraphs[0].add_run("Page No.")

    # Exact page mapping according to controlled section pagination
    index_data = [
        ("1", "Introduction", "2"),
        ("2", "Project Purpose and Objectives", "2"),
        ("3", "Technological Stack", "3"),
        ("4", "Database Architecture", "4"),
        ("5", "Application Workflow and Screenshots", "5"),
        ("", "5.1 Authentication and Account Setup", "5"),
        ("", "5.2 Student Portal", "8"),
        ("", "5.3 Company Portal", "13"),
        ("", "5.4 Placement Officer (Admin) Portal", "19"),
        ("", "5.5 Automated Email Notification System", "23"),
        ("6", "Conclusion", "26"),
    ]

    for sr, topic, pg in index_data:
        row_cells = index_table.add_row().cells
        row_cells[0].paragraphs[0].add_run(sr)
        row_cells[1].paragraphs[0].add_run(topic)
        row_cells[2].paragraphs[0].add_run(pg)

    format_table_cells(index_table, [Inches(1.0), Inches(4.2), Inches(1.3)])
    
    doc.add_page_break()

    # =========================================================================
    # PAGE 2: 1. INTRODUCTION & 2. PROJECT PURPOSE AND OBJECTIVES
    # =========================================================================
    add_heading_1("1. Introduction")
    add_body_p(
        "Campus recruitment is one of the most critical events in any college or university. Every year, hundreds of students participate in placement drives organized by visiting companies. However, in most institutions, the process is still handled manually using disconnected tools—such as Google Forms, Excel sheets, email notices, and WhatsApp groups."
    )
    add_body_p(
        "This traditional approach creates serious challenges: students often miss drive notices, placement officers spend days verifying eligibility criteria like CGPA and backlogs by hand, and recruiters face difficulty shortlisting candidates efficiently."
    )
    add_body_p(
        "PlaceMate is a comprehensive web application built to solve these exact problems. It automates the entire placement lifecycle by providing three tailored portals: one for Students, one for Companies, and one for College Placement Officers (Admin). The platform centralizes student profiles, automates eligibility checks, manages drive applications, delivers branded email alerts, and provides visual recruitment analytics."
    )

    add_heading_1("2. Project Purpose and Objectives")
    
    add_heading_2("Project Purpose")
    add_body_p(
        "The primary purpose of PlaceMate is to replace chaotic, manual placement operations with a modern, transparent, and automated digital system. By providing a single source of truth, PlaceMate ensures that students only see drives they are actually eligible for, companies can post requirements and manage applicants seamlessly, and placement officers have real-time visibility into hiring outcomes."
    )

    add_heading_2("Core Objectives")
    add_bullet(" To provide secure, distinct dashboards for Students, Company HRs, and Placement Officers with granular role-based permissions.", "Role-Based Portals:")
    add_bullet(" To verify email addresses during registration using a time-limited 6-digit OTP sent via SMTP, preventing fake user registrations.", "OTP-Verified Authentication:")
    add_bullet(" To automatically validate student criteria (CGPA, active backlogs, passing year, and department) against company drive criteria before an application can be submitted.", "Automated Eligibility Filtering:")
    add_bullet(" To give students a transparent view of their job applications (Applied, Shortlisted, Interview Scheduled, Selected, Rejected) with instant status updates.", "Real-Time Application Tracking:")
    add_bullet(" To automatically send professional, responsive HTML email notifications whenever a student registers, applies for a drive, or gets shortlisted/selected.", "Automated Email Notification System:")
    add_bullet(" To display visual charts (department hiring rates, package distribution, application ratios) using Chart.js for both students and company recruiters.", "Interactive Analytics:")
    add_bullet(" To deliver a responsive, accessible interface with zero-flicker light and dark mode toggles integrated across every screen.", "Modern UI and Theming:")

    doc.add_page_break()

    # =========================================================================
    # PAGE 3: 3. TECHNOLOGICAL STACK
    # =========================================================================
    add_heading_1("3. Technological Stack")
    add_body_p(
        "PlaceMate is developed using a robust, production-ready stack chosen for performance, security, and developer productivity:"
    )

    tech_table = doc.add_table(rows=1, cols=3)
    tech_hdr = tech_table.rows[0].cells
    tech_hdr[0].paragraphs[0].add_run("Component")
    tech_hdr[1].paragraphs[0].add_run("Technology")
    tech_hdr[2].paragraphs[0].add_run("Role in Project")

    tech_data = [
        ("Backend Framework", "Python 3.13 & Django", "Handles application routing, business logic, ORM data queries, user authentication, and security protections (CSRF, XSS, password hashing)."),
        ("Database", "PostgreSQL (Supabase)", "Stores all relational data including user accounts, profiles, drives, applications, and master records with strict foreign keys."),
        ("Frontend & Styling", "HTML5 & Vanilla CSS", "Built with custom CSS variables, responsive flexbox and grid layouts, clean cards, and unified light/dark color themes without heavy dependencies."),
        ("Client Scripting", "Vanilla JavaScript", "Controls client-side OTP timers, instant form validation, dynamic theme toggling, and modal interactions."),
        ("Data Visualization", "Chart.js", "Generates interactive doughnut and bar charts for student placement analytics and company recruitment metrics."),
        ("Email Service", "Django SMTP & Gmail", "Delivers responsive HTML emails for OTP codes, application confirmations, status updates, and password resets."),
        ("Static Assets", "WhiteNoise", "Handles efficient serving and caching of static stylesheets, scripts, and logos in production."),
    ]

    for comp, tech, role in tech_data:
        r_cells = tech_table.add_row().cells
        r_cells[0].paragraphs[0].add_run(comp)
        r_cells[1].paragraphs[0].add_run(tech)
        r_cells[2].paragraphs[0].add_run(role)

    format_table_cells(tech_table, [Inches(1.8), Inches(1.8), Inches(2.9)])

    doc.add_page_break()

    # =========================================================================
    # PAGE 4: 4. DATABASE ARCHITECTURE
    # =========================================================================
    add_heading_1("4. Database Architecture")
    add_body_p(
        "PlaceMate relies on a relational database schema designed to maintain clean separation of concerns while enforcing referential integrity. The database is organized into distinct entities mapped through Django's Object-Relational Mapping (ORM)."
    )

    add_heading_2("Core Entities")
    add_bullet(" Extends Django's AbstractUser. Uses email as the unique username identifier and includes a role field ('student', 'company', 'officer') along with an email_verified boolean flag.", "User (Custom User Model):")
    add_bullet(" Linked 1-to-1 with User. Stores student academic details including full name, roll number, university, department (Foreign Key), semester, CGPA, passing year, active backlogs, verified status, resume file, and skills (Many-to-Many).", "StudentProfile:")
    add_bullet(" Linked 1-to-1 with User. Stores corporate profile details including company name, logo image, industry type, website URL, headquarter location, description, and admin verification status.", "CompanyProfile:")
    add_bullet(" Linked 1-to-1 with User. Represents college placement coordinators who verify drives and oversee campus hiring activities.", "OfficerProfile:")
    add_bullet(" Master lookup tables storing academic departments (e.g., Computer Engineering, Information Technology) and technical skills (e.g., Python, React, SQL).", "Department and Skill:")
    add_bullet(" Created by CompanyProfile. Specifies job title, description, package (CTC in LPA), job location, eligible departments (Many-to-Many), required skills (Many-to-Many), minimum CGPA cutoff, maximum allowed backlogs, eligible passing year, deadlines, and drive status ('draft', 'published', 'closed').", "PlacementDrive:")
    add_bullet(" Bridges StudentProfile and PlacementDrive (Many-to-Many junction). Tracks candidate application status ('applied', 'shortlisted', 'interview_scheduled', 'selected', 'rejected') and maintains application timestamps.", "Application:")

    add_heading_2("Entity Relationship Overview")
    
    er_table = doc.add_table(rows=1, cols=4)
    er_hdr = er_table.rows[0].cells
    er_hdr[0].paragraphs[0].add_run("Source Entity")
    er_hdr[1].paragraphs[0].add_run("Relationship")
    er_hdr[2].paragraphs[0].add_run("Target Entity")
    er_hdr[3].paragraphs[0].add_run("Description")

    er_data = [
        ("User", "One-to-One", "StudentProfile", "Each student account maps to exactly one student profile."),
        ("User", "One-to-One", "CompanyProfile", "Each recruiter account maps to exactly one company profile."),
        ("User", "One-to-One", "OfficerProfile", "Each staff account maps to a placement officer profile."),
        ("StudentProfile", "Many-to-One", "Department", "Every student belongs to an academic department."),
        ("StudentProfile", "Many-to-Many", "Skill", "Students can add multiple skills to their profile."),
        ("CompanyProfile", "One-to-Many", "PlacementDrive", "A company can publish multiple placement drives."),
        ("PlacementDrive", "Many-to-Many", "Department", "A drive can be open to multiple eligible departments."),
        ("StudentProfile", "One-to-Many", "Application", "A student can apply to multiple eligible drives."),
        ("PlacementDrive", "One-to-Many", "Application", "A drive receives applications from multiple students."),
    ]

    for src, rel, tgt, desc in er_data:
        r_cells = er_table.add_row().cells
        r_cells[0].paragraphs[0].add_run(src)
        r_cells[1].paragraphs[0].add_run(rel)
        r_cells[2].paragraphs[0].add_run(tgt)
        r_cells[3].paragraphs[0].add_run(desc)

    format_table_cells(er_table, [Inches(1.5), Inches(1.2), Inches(1.5), Inches(2.3)])

    doc.add_page_break()

    # =========================================================================
    # PAGE 5: 5. APPLICATION WORKFLOW AND SCREENSHOTS
    # 5.1 Authentication and Account Setup (Part 1)
    # =========================================================================
    add_heading_1("5. Application Workflow and Screenshots")
    add_body_p(
        "This section presents the step-by-step visual workflow of PlaceMate, walking through Authentication, the Student Portal, the Company Portal, the Placement Officer Portal, and the Automated Email System."
    )

    add_heading_2("5.1 Authentication and Account Setup")
    add_body_p(
        "PlaceMate provides a secure, role-specific onboarding experience. Users select their appropriate role, verify their official email address using a one-time password (OTP), and log in securely."
    )

    fig_count = 1

    add_screenshot(
        "Screenshots/Registration_1.png",
        fig_count, "Registration Role Selection Screen",
        "The starting onboarding screen where visitors select whether they want to register as a Student seeking job opportunities or as a Company looking to hire fresh talent."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Student Side/Student_Registration.png",
        fig_count, "Student Registration Form",
        "The student signup screen where candidates enter their email, create a secure password, and fill in academic credentials including roll number, department, CGPA, and passing year."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 6: 5.1 Authentication (Part 2)
    # =========================================================================
    add_screenshot(
        "Screenshots/OTP.png",
        fig_count, "Real-Time Email OTP Verification",
        "The integrated OTP verification box that activates upon clicking 'Send OTP'. It prevents fraudulent registrations by ensuring students possess access to their registered email inbox."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Company Siide/Company_Registration.png",
        fig_count, "Company Registration Form",
        "The recruiter registration screen allowing HR representatives to register their corporate profile with email verification, company name, industry, website, and headquarters location."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 7: 5.1 Authentication (Part 3)
    # =========================================================================
    add_screenshot(
        "Screenshots/Login.png",
        fig_count, "PlaceMate Central Login Screen",
        "The unified login screen featuring an integrated dark/light mode toggle in the top bar. Users enter their credentials and are automatically routed to their respective role-based portal."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Reset_Password.png",
        fig_count, "Password Reset Request Screen",
        "Self-service password recovery screen where users enter their registered email address to receive a secure, time-limited password reset link."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Set_New_Password.png",
        fig_count, "Set New Password Screen",
        "Secure token-authenticated page allowing the user to create a new password adhering to complexity requirements.",
        width_in=4.4
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 8: 5.2 Student Portal (Part 1)
    # =========================================================================
    add_heading_2("5.2 Student Portal")
    add_body_p(
        "The Student Portal serves as a personalized career hub where candidates can maintain their academic profile, discover placement drives, check eligibility criteria in real time, submit job applications, and analyze campus placement trends."
    )

    add_screenshot(
        "Screenshots/Student Side/Dashboard.png",
        fig_count, "Student Dashboard",
        "The main student dashboard displaying quick summary cards (eligible drives, active applications, shortlisted status) and a feed of newly announced placement drives."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Student Side/Profile.png",
        fig_count, "Student Profile Overview",
        "Detailed view of the student's academic profile showing roll number, department, CGPA, active backlogs, verified status badge, tagged technical skills, and uploaded resume link."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 9: 5.2 Student Portal (Part 2)
    # =========================================================================
    add_screenshot(
        "Screenshots/Student Side/Edit_Profile.png",
        fig_count, "Edit Student Profile",
        "Profile editing interface where students can update their skills, upload new resumes, and refresh academic details."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Student Side/Placement_Drive.png",
        fig_count, "Placement Drives Listing",
        "List of all published recruitment drives on campus with search filters, salary tags, deadline alerts, and quick eligibility badges."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 10: 5.2 Student Portal (Part 3)
    # =========================================================================
    add_screenshot(
        "Screenshots/Student Side/Drive_1.png",
        fig_count, "Placement Drive Details & Eligibility Breakdown",
        "In-depth view of a specific placement drive outlining job role, CTC, location, required skill tags, application deadline, and automated eligibility criteria checks."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Student Side/Drive_2.png",
        fig_count, "Drive Application Submission",
        "Drive application confirmation modal that allows students who meet all criteria to submit their resume and application with a single click."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 11: 5.2 Student Portal (Part 4)
    # =========================================================================
    add_screenshot(
        "Screenshots/Student Side/Applications.png",
        fig_count, "My Applications Tracking Screen",
        "Dedicated status board where students track every submitted application in real time, complete with color-coded status badges (Applied, Shortlisted, Selected)."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Student Side/Analytics_1.png",
        fig_count, "Student Placement Analytics - Key Performance Indicators",
        "Analytics dashboard providing students with clear insights on total companies visited, placement percentage, and overall drive participation."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 12: 5.2 Student Portal (Part 5)
    # =========================================================================
    add_screenshot(
        "Screenshots/Student Side/Analytics_2.png",
        fig_count, "Student Analytics - Department Hiring Trends",
        "Interactive Chart.js chart illustrating placement counts and selection rates broken down across engineering departments."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Student Side/Analytics_3.png",
        fig_count, "Student Analytics - Salary Package Distribution",
        "Visualization showing the distribution of salary packages (CTC in LPA) across campus placement drives."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 13: 5.3 Company Portal (Part 1)
    # =========================================================================
    add_heading_2("5.3 Company Portal")
    add_body_p(
        "The Company Portal empowers corporate recruiters to manage their company profile, create customized placement drives with specific academic cutoffs, review student applicants, and manage candidate selection stages."
    )

    add_screenshot(
        "Screenshots/Company Siide/Dashboard.png",
        fig_count, "Company Recruiter Dashboard",
        "The recruiter dashboard showing high-level recruitment metrics, including total drives posted, total applicants received, and shortlisted candidates."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Company Siide/Profile.png",
        fig_count, "Company Profile Overview",
        "Corporate profile page displaying company branding, logo, industry sector, website, office location, and verification status."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 14: 5.3 Company Portal (Part 2)
    # =========================================================================
    add_screenshot(
        "Screenshots/Company Siide/Edit Profile.png",
        fig_count, "Edit Company Profile",
        "Form allowing HR personnel to update recruiter contact information, company bio, and recruitment preferences."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Company Siide/Placement_Drive.png",
        fig_count, "Company Placement Drives Management",
        "Overview table of all placement drives published by the company, showing application deadlines, applicant counts, and drive statuses."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 15: 5.3 Company Portal (Part 3)
    # =========================================================================
    add_screenshot(
        "Screenshots/Company Siide/New_Drive_1.png",
        fig_count, "Create Placement Drive - Job Role and Details",
        "Step 1 of drive creation where recruiters enter the job title, job description, annual CTC package, job location, and application deadline."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Company Siide/New_Drive_2.png",
        fig_count, "Create Placement Drive - Eligibility Criteria",
        "Step 2 of drive creation where recruiters configure strict eligibility rules: minimum CGPA cutoff, maximum backlogs, eligible passing year, and required technical skills."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 16: 5.3 Company Portal (Part 4)
    # =========================================================================
    add_screenshot(
        "Screenshots/Company Siide/Applicants.png",
        fig_count, "Drive Applicants Management Screen",
        "Applicant review table listing all students who applied to the drive, complete with roll number, department, CGPA, resume download link, and current stage."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Company Siide/Placed_Drive_1.png",
        fig_count, "Applicant Status Review and Updating",
        "Interface allowing company HR to review candidate qualifications and update their application status to Shortlisted, Interview Scheduled, or Rejected."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 17: 5.3 Company Portal (Part 5)
    # =========================================================================
    add_screenshot(
        "Screenshots/Company Siide/Placed_Drive_2.png",
        fig_count, "Final Selection and Offer Confirmation",
        "Candidate management view confirming final student selections and triggering automated congratulatory emails."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Company Siide/Analytics_1.png",
        fig_count, "Company Analytics - Hiring Funnel",
        "Recruitment analytics screen presenting the company's application-to-selection funnel and average applicant CGPA."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 18: 5.3 Company Portal (Part 6)
    # =========================================================================
    add_screenshot(
        "Screenshots/Company Siide/Analytics_2.png",
        fig_count, "Company Analytics - Department-wise Candidate Breakdown",
        "Visual chart detailing applicant distribution by academic branch to help recruiters evaluate candidate pool diversity."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 19: 5.4 Placement Officer (Admin) Portal (Part 1)
    # =========================================================================
    add_heading_2("5.4 Placement Officer (Admin) Portal")
    add_body_p(
        "The Placement Officer Portal provides college administrators with centralized control over users, verification requests, master academic records, drives, and institutional hiring statistics."
    )

    add_screenshot(
        "Screenshots/Admin/Dashboard.png",
        fig_count, "Placement Officer Admin Dashboard",
        "Central administration dashboard showing university-wide recruitment metrics, pending company verifications, and active campus drives."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Admin/Users.png",
        fig_count, "User Management and Verification",
        "User administration screen where placement officers verify student and company accounts, assign roles, and manage system access."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 20: 5.4 Placement Officer Portal (Part 2)
    # =========================================================================
    add_screenshot(
        "Screenshots/Admin/Depart,ments.png",
        fig_count, "Academic Department Master Management",
        "Master data screen allowing administrators to configure college departments, degree codes, and eligible streams."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Admin/Skills.png",
        fig_count, "Technical Skills Master Management",
        "Master catalogue of technical skills and competencies used for student tagging and drive matching."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 21: 5.4 Placement Officer Portal (Part 3)
    # =========================================================================
    add_screenshot(
        "Screenshots/Admin/Placement Drives.png",
        fig_count, "Placement Drives Oversight and Approval",
        "Administration screen for reviewing drives submitted by visiting companies, confirming eligibility criteria, and approving them for publishing."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Admin/Application.png",
        fig_count, "Global Application Records Oversight",
        "Comprehensive database log of all student applications across all active and completed campus placement drives."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 22: 5.4 Placement Officer Portal (Part 4)
    # =========================================================================
    add_screenshot(
        "Screenshots/Admin/Other_profiles.png",
        fig_count, "Placement Officer Staff Profiles",
        "Management interface for placement cell faculty members and administrative staff coordinators."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 23: 5.5 Automated Email Notification System (Part 1)
    # =========================================================================
    add_heading_2("5.5 Automated Email Notification System")
    add_body_p(
        "PlaceMate incorporates an automated email notification system built with custom, responsive HTML templates. Whenever significant actions take place, the server dispatches instant, branded notifications directly to the user's inbox."
    )

    add_screenshot(
        "Screenshots/Mails/Registered Mail.png",
        fig_count, "Welcome & Account Verification Email",
        "Automated welcome email dispatched upon successful registration, welcoming the user to PlaceMate and confirming their verified credentials."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Mails/Application_Mail.png",
        fig_count, "Drive Application Confirmation Email",
        "Confirmation email received by students immediately upon applying for a drive, summarizing the job title, company name, and application reference."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 24: 5.5 Automated Email Notification System (Part 2)
    # =========================================================================
    add_screenshot(
        "Screenshots/Mails/Status_Shotlisted.png",
        fig_count, "Candidate Shortlisted Notification Email",
        "Real-time alert sent to candidates whose applications have been reviewed and shortlisted for the upcoming interview round."
    )
    fig_count += 1

    add_screenshot(
        "Screenshots/Mails/Status_Selected.png",
        fig_count, "Final Selection & Job Offer Congratulatory Email",
        "Official congratulatory email sent when a company marks a candidate as Selected, providing job confirmation and package details."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 25: 5.5 Automated Email Notification System (Part 3)
    # =========================================================================
    add_screenshot(
        "Screenshots/Mails/Reset_Password_Mail.png",
        fig_count, "Secure Password Reset Email",
        "Password recovery email containing a secure, single-use reset link with clear security advice."
    )
    fig_count += 1

    doc.add_page_break()

    # =========================================================================
    # PAGE 26: 6. CONCLUSION
    # =========================================================================
    add_heading_1("6. Conclusion")
    add_body_p(
        "PlaceMate successfully transforms the traditional, fragmented campus placement process into a cohesive, automated digital platform. By uniting Students, Recruiters, and College Placement Officers within a single system, it eliminates manual data verification, prevents human errors in eligibility checking, and ensures transparent communication through real-time application tracking and automated emails."
    )
    add_body_p(
        "The application's responsive design, integrated dark and light mode system, and interactive analytics dashboards provide an intuitive and efficient experience for all users. PlaceMate delivers a scalable, production-ready solution that modernizes institutional recruitment and helps colleges run smooth, organized, and data-driven placement seasons."
    )

    # Output file
    output_filename = "PlaceMate_Project_Documentation.docx"
    doc.save(output_filename)
    print(f"Documentation successfully created: {output_filename}")

if __name__ == "__main__":
    create_placemate_doc()
